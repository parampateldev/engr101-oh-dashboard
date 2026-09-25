#!/usr/bin/env python3
"""Convert the ENGR101 office hours Excel schedule into schedule.json."""

import json
import sys
from datetime import time
from pathlib import Path

try:
    import openpyxl
except ImportError:
    raise SystemExit("Install openpyxl: pip install openpyxl")

DEFAULT_XLSX = Path.home() / "Downloads/F26 Lab and Office Hours Schedule + Attendance.xlsx"
OUT = Path(__file__).resolve().parent / "schedule.json"
DAYS = ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]
OH_COLS = range(10, 16)
TOTAL_COL = 16


import re


def format_display_name(raw_name: str) -> str:
    """Format 'Last, First Middle' or 'First Last' into 'First L.' (first name and first initial of last name)."""
    if not raw_name or not isinstance(raw_name, str):
        return ""
    cleaned = re.sub(r"\(.*?\)", "", raw_name).strip()
    if not cleaned:
        return ""
    if "," in cleaned:
        parts = [p.strip() for p in cleaned.split(",", 1)]
        last = parts[0].strip()
        first_tokens = parts[1].split()
        first = first_tokens[0] if first_tokens else ""
        last_initial = f"{last[0].upper()}." if last else ""
        return f"{first} {last_initial}".strip()
    tokens = cleaned.split()
    if len(tokens) == 1:
        return tokens[0]
    first = tokens[0]
    last = tokens[-1].rstrip(".")
    last_initial = f"{last[0].upper()}." if last else ""
    return f"{first} {last_initial}".strip()


def parse_staff_names(wb) -> dict:
    if "Overview" not in wb.sheetnames:
        return {}

    ws = wb["Overview"]
    names = {}
    for row in range(1, ws.max_row + 1):
        full_name = ws.cell(row, 1).value
        uniqname = ws.cell(row, 2).value
        if (
            not full_name
            or not uniqname
            or not isinstance(full_name, str)
            or not isinstance(uniqname, str)
            or full_name.strip() in {"Staff", "Uniqname"}
        ):
            continue
        names[uniqname.strip()] = full_name.strip()
    return names


def parse_student_names(wb) -> dict:
    sheet_name = next((s for s in wb.sheetnames if "student" in s.lower() and "roster" in s.lower()), None)
    if not sheet_name:
        return {}
    ws = wb[sheet_name]
    students = {}
    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 2:
            continue
        name, uniq = row[0], row[1]
        if (
            not name
            or not uniq
            or not isinstance(name, str)
            or not isinstance(uniq, str)
            or uniq.strip() in {"SIS Login ID", "Uniqname", "uniqname"}
        ):
            continue
        fmt = format_display_name(name)
        if fmt:
            students[uniq.strip().lower()] = fmt
    return students


def parse_schedule(xlsx_path: Path) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    staff_names = parse_staff_names(wb)
    student_names = parse_student_names(wb)
    ws = wb["The ScheduleTM"]
    schedule = {}

    for day in DAYS:
        start = next((r + 1 for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == day), None)
        if not start:
            continue

        end = ws.max_row + 1
        for r in range(start, ws.max_row + 1):
            if ws.cell(r, 1).value in DAYS:
                end = r
                break

        slots = []
        for r in range(start, end):
            slot_time = ws.cell(r, 1).value
            if not isinstance(slot_time, time):
                continue

            staff = []
            for col in OH_COLS:
                value = ws.cell(r, col).value
                if (
                    value
                    and isinstance(value, str)
                    and value.strip()
                    and "Slot" not in value
                    and value != "STAFF MEETING"
                ):
                    staff.append(value.strip())

            total = ws.cell(r, TOTAL_COL).value
            if staff or total:
                slots.append(
                    {
                        "time": slot_time.strftime("%H:%M"),
                        "staff": staff,
                        "total": int(total) if total else len(staff),
                    }
                )

        schedule[day.lower()] = slots

    return {
        "source": xlsx_path.name,
        "timezone": "America/Detroit",
        "staff_names": staff_names,
        "student_names": student_names,
        "schedule": schedule,
    }


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx.is_file():
        raise SystemExit(f"Schedule file not found: {xlsx}")

    OUT.write_text(json.dumps(parse_schedule(xlsx), indent=2))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
